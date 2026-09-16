"""
Insight-card -> PPTX export, matching the ProcDNA-style branded template
(gradient top bar, numbered badge + title header, "key takeaways" summary
strip, a card-type-specific main panel, and a WHAT THIS MEANS / SOURCES
footer row) instead of the earlier plain white .icard look.

Layout model
------------
The slide is a FIXED standard 16:9 (13.333" x 7.5"). Header, summary strip,
bottom row, and footer all size themselves from their own (usually short,
predictable) text. Whatever vertical space is left in between is the
"content area", and the evidence renderer for the card's type (metrics /
table / steps / synthesis grid) is fit into that fixed space with the same
binary-search "shrink fonts & padding until it fits" approach used
everywhere else in this file — never by growing the slide.

Text is measured with the actual Carlito glyph metrics (metric-compatible
with Calibri, i.e. what LibreOffice/Word really render) rather than a
character-count guess, via `_text_width_in`.
"""

import io
import re

from fastapi import HTTPException
from fastapi.responses import StreamingResponse
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.util import Emu, Inches, Pt

# --------------------------------------------------------------------------- #
# Palette — sampled from the reference template
# --------------------------------------------------------------------------- #
C_NAVY = RGBColor(0x1B, 0x2A, 0x6E)          # titles, metric values
C_BLUE = RGBColor(0x2F, 0x5F, 0xDE)          # primary accent (badges, headings, pills)
C_BLUE_LIGHT_BG = RGBColor(0xEA, 0xF1, 0xFE)  # badge / pill / what-this-means fill
C_SUMMARY_BG = RGBColor(0xEE, 0xF2, 0xFC)
C_MUTED_NAVY = RGBColor(0x47, 0x55, 0x6B)
C_TABLE_HEADER_BG = RGBColor(0xEE, 0xF1, 0xF6)
C_TABLE_TEXT = RGBColor(0x1E, 0x29, 0x3B)
C_TABLE_FAINT = RGBColor(0x8A, 0x94, 0xA6)
C_BORDER_LIGHT = RGBColor(0xDD, 0xE3, 0xEE)
C_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
C_TEAL = RGBColor(0x12, 0xB5, 0xA6)
C_GREEN_ICON = RGBColor(0x12, 0xA1, 0x87)
C_PURPLE_ICON = RGBColor(0x7C, 0x5C, 0xFC)
C_ICON_BG = [
    (RGBColor(0xE4, 0xEC, 0xFB), C_BLUE),
    (RGBColor(0xE1, 0xF5, 0xEE), C_GREEN_ICON),
    (RGBColor(0xEE, 0xE9, 0xFB), C_PURPLE_ICON),
    (RGBColor(0xDF, 0xF7, 0xF1), C_TEAL),
]
C_PASTEL = [
    RGBColor(0xE7, 0xEC, 0xFC), RGBColor(0xE3, 0xF5, 0xEC),
    RGBColor(0xEF, 0xEA, 0xFC), RGBColor(0xFB, 0xEF, 0xE3),
]
C_FOOTER_GRAY = RGBColor(0x9A, 0xA3, 0xB2)
GRADIENT_LEFT = RGBColor(0x74, 0xB4, 0xF2)
GRADIENT_RIGHT = RGBColor(0x1E, 0x4E, 0xD8)

FONT = "Calibri"

SLIDE_W_IN = 13.333
SLIDE_H_IN = 7.5
MARGIN_X_IN = 0.45
MIN_SCALE = 0.55
MAX_SCALE = 1.35

_ACRONYMS = {
    "nci", "seer", "pmc", "ash", "esmo", "fda", "who", "mrd", "ctgov",
    "clinicaltrials", "ppt", "ash/blood",
}


# --------------------------------------------------------------------------- #
# Text measurement — real glyph widths via Carlito (metric-compatible with
# Calibri), falling back to a calibrated character-count estimate if the font
# file isn't present on the host.
# --------------------------------------------------------------------------- #
_FONT_FILES = {
    (False, False): "/usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf",
    (True, False): "/usr/share/fonts/truetype/crosextra/Carlito-Bold.ttf",
    (False, True): "/usr/share/fonts/truetype/crosextra/Carlito-Italic.ttf",
    (True, True): "/usr/share/fonts/truetype/crosextra/Carlito-BoldItalic.ttf",
}
_FONT_REF_SIZE = 200
_font_cache = {}


def _get_font(bold=False, italic=False):
    key = (bold, italic)
    if key in _font_cache:
        return _font_cache[key]
    try:
        from PIL import ImageFont
        font = ImageFont.truetype(_FONT_FILES[key], _FONT_REF_SIZE)
    except Exception:
        font = None
    _font_cache[key] = font
    return font


def _char_w(bold=False):
    return 0.47 if bold else 0.43


def _text_width_in(text, font_pt, bold=False, italic=False):
    if not text:
        return 0.0
    font = _get_font(bold, italic)
    if font is None:
        return (len(text) * font_pt * _char_w(bold)) / 72.0
    return (font.getlength(text) * (font_pt / _FONT_REF_SIZE)) / 72.0


def wrap_text(text, width_in, font_pt, bold=False):
    text = (text or "").strip()
    if not text:
        return [""]
    width_in = max(width_in, 0.3)
    font = _get_font(bold)
    if font is None:
        import textwrap
        chars_per_line = max(6, int((width_in * 72) / (font_pt * _char_w(bold))))
        lines = []
        for para in text.split("\n"):
            lines.extend(textwrap.wrap(para, width=chars_per_line) or [""])
        return lines
    lines = []
    for para in text.split("\n"):
        words = para.split(" ")
        cur = ""
        for word in words:
            trial = f"{cur} {word}".strip()
            if not cur or _text_width_in(trial, font_pt, bold) <= width_in:
                cur = trial
            else:
                lines.append(cur)
                cur = word
        lines.append(cur)
    return lines or [""]


def text_height(text, width_in, font_pt, bold=False, line_mult=1.24):
    n = len(wrap_text(text, width_in, font_pt, bold))
    return (n * font_pt * line_mult) / 72.0


def pill_w(text, font_pt, pad_in=0.30, bold=True, min_in=0.55):
    return max(min_in, pad_in + _text_width_in(text, font_pt, bold))


def clean_text(txt):
    return re.sub(r"<[^>]+>", "", str(txt)).replace("[VERIFIED]", "").replace("[INFERENCE]", "").strip()


def humanize_source_id(sid):
    words = re.split(r"[_\-]+", sid.strip())
    out = [w.upper() if w.lower() in _ACRONYMS else w.capitalize() for w in words]
    return " ".join(out) or sid


def resolve_source_name(store, run_id, sid, source_names=None):
    if source_names and sid in source_names:
        return source_names[sid]
    getter = getattr(store, "get_source_display_name", None)
    if callable(getter):
        try:
            name = getter(run_id, sid)
            if name:
                return name
        except Exception:
            pass
    return humanize_source_id(sid)


def pick_icon(text, idx=0):
    t = (text or "").lower()
    rules = [
        (("incidence",), "person"),
        (("mortality", "death"), "heart"),
        (("living", "population", "prevalent", "estimated"), "people"),
        (("survival", "rate", "percent", "%"), "bars"),
        (("demographic", "common in", "segment"), "people"),
        (("presentation", "symptom", "clinical"), "person"),
        (("chemotherapy", "treatment", "induction", "therapy", "regimen", "frontline"), "drip"),
        (("remission", "complete"), "check"),
        (("mrd", "chromosome", "genetic", "molecular"), "dna"),
        (("risk", "response", "assessment"), "bars"),
        (("relapse",), "refresh"),
        (("transplant", "stem cell"), "people"),
    ]
    for keywords, kind in rules:
        if any(k in t for k in keywords):
            return kind
    fallback = ["person", "heart", "people", "bars", "dna", "check", "refresh", "document"]
    return fallback[idx % len(fallback)]


# --------------------------------------------------------------------------- #
# Drawing primitives — every one accepts `slide=None` and skips drawing, so
# the same code path measures (pass 1) and draws (pass 2).
# --------------------------------------------------------------------------- #
def add_soft_shadow(shape, blur_emu=70000, dist_emu=18000, alpha=18000, color="17223A"):
    spPr = shape._element.spPr
    for existing in spPr.findall(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}effectLst"
    ):
        spPr.remove(existing)
    xml = (
        '<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f'<a:outerShdw blurRad="{blur_emu}" dist="{dist_emu}" dir="5400000" rotWithShape="0">'
        f'<a:srgbClr val="{color}"><a:alpha val="{alpha}"/></a:srgbClr>'
        "</a:outerShdw></a:effectLst>"
    )
    spPr.append(parse_xml(xml))


def _clear_effects(shape):
    spPr = shape._element.spPr
    spPr.append(parse_xml(
        '<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>'
    ))


def add_shape(slide, kind, x, y, w, h, fill=None, line=None, line_w=1.0, adjustments=None):
    if slide is None:
        return None
    shp = slide.shapes.add_shape(kind, x, y, w, h)
    if adjustments:
        try:
            for i, v in enumerate(adjustments):
                shp.adjustments[i] = v
        except Exception:
            pass
    if fill is not None:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    else:
        shp.fill.background()
    if line is not None:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    _clear_effects(shp)
    return shp


def add_rect(slide, x, y, w, h, fill=None, line=None, radius=None, shadow=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    shp = add_shape(slide, kind, x, y, w, h, fill=fill, line=line, adjustments=[radius] if radius is not None else None)
    if shp is not None and shadow:
        add_soft_shadow(shp)
    return shp


def add_gradient_bar(slide, x, y, w, h, left_color, right_color):
    if slide is None:
        return
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.line.fill.background()
    shp.shadow.inherit = False
    fill = shp.fill
    fill.gradient()
    stops = fill.gradient_stops
    stops[0].position, stops[0].color.rgb = 0.0, left_color
    stops[1].position, stops[1].color.rgb = 1.0, right_color
    try:
        fill.gradient_angle = 0.0  # left -> right
    except Exception:
        pass


def add_text(slide, x, y, w, h, text, font_pt, color, bold=False, italic=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_mult=1.24, font_name=FONT):
    """Pre-wraps with the exact same `wrap_text()` used for height measurement,
    then disables the textbox's own word-wrap. This guarantees the rendered
    line breaks always match what was measured — and, critically, means a
    single word that's wider than the box overflows visibly instead of being
    split mid-word by the renderer's own fallback wrapping."""
    if slide is None:
        return None
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    width_in = max(w / 914400, 0.05)
    first = True
    for para in (text or "").split("\n"):
        for line in wrap_text(para, width_in, font_pt, bold):
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.text = line
            p.alignment = align
            p.line_spacing = line_mult
            r = p.runs[0]
            r.font.size = Pt(font_pt)
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = color
            r.font.name = font_name
    return box


def add_pill(slide, x, y, w, h, text, fill, text_color, font_pt=10, bold=True,
             border=None, align=PP_ALIGN.CENTER, font_name=FONT):
    shp = add_rect(slide, x, y, w, h, fill=fill, line=border, radius=0.5)
    if slide is None:
        return None
    tf = shp.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    r = p.runs[0]
    r.font.size = Pt(font_pt)
    r.font.bold = bold
    r.font.color.rgb = text_color
    r.font.name = font_name
    return shp


# -- icon glyphs -------------------------------------------------------------- #
def draw_icon(slide, kind, x, y, d, color):
    """Draws a simple abstract pictogram inside the d x d box at (x, y)."""
    if slide is None:
        return
    if kind == "person":
        _icon_person(slide, x, y, d, color)
    elif kind == "people":
        _icon_person(slide, x, y, d, color)
    elif kind == "heart":
        add_shape(slide, MSO_SHAPE.HEART, x, y + d * 0.06, d, d * 0.86, fill=color)
    elif kind == "bars":
        bw = d * 0.24
        gap = d * 0.10
        heights = [d * 0.45, d * 0.68, d * 0.9]
        cx = x
        for hgt in heights:
            add_shape(slide, MSO_SHAPE.RECTANGLE, cx, y + d - hgt, bw, hgt, fill=color)
            cx += bw + gap
    elif kind == "dna":
        add_shape(slide, MSO_SHAPE.WAVE, x, y + d * 0.32, d, d * 0.36, fill=color)
    elif kind == "drip":
        try:
            add_shape(slide, MSO_SHAPE.TEAR, x + d * 0.2, y, d * 0.6, d * 0.75, fill=color)
        except Exception:
            add_shape(slide, MSO_SHAPE.OVAL, x + d * 0.2, y + d * 0.25, d * 0.6, d * 0.6, fill=color)
        add_shape(slide, MSO_SHAPE.RECTANGLE, x + d * 0.46, y + d * 0.75, d * 0.08, d * 0.2, fill=color)
    elif kind == "document":
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x + d * 0.15, y, d * 0.7, d,
                  fill=None, line=color, line_w=1.5, adjustments=[0.12])
        for i, frac in enumerate([0.3, 0.5, 0.7]):
            add_shape(slide, MSO_SHAPE.RECTANGLE, x + d * 0.28, y + d * frac, d * 0.44, d * 0.05, fill=color)
    elif kind in ("check", "refresh"):
        glyph = "\u2713" if kind == "check" else "\u21bb"
        d_in = d / 914400
        add_text(slide, x, y, d, d, glyph, d_in * 54, color, bold=True,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    else:
        add_shape(slide, MSO_SHAPE.OVAL, x, y, d, d, fill=color)


def _icon_person(slide, x, y, d, color, alpha_fill=False):
    head_d = d * 0.42
    add_shape(slide, MSO_SHAPE.OVAL, x + (d - head_d) / 2, y, head_d, head_d, fill=color)
    body_w, body_h = d * 0.82, d * 0.5
    add_shape(slide, MSO_SHAPE.TRAPEZOID, x + (d - body_w) / 2, y + head_d * 0.82, body_w, body_h, fill=color)


# --------------------------------------------------------------------------- #
# Header, summary strip, bottom row, footer — fixed-ish sections
# --------------------------------------------------------------------------- #
def render_top_bar(slide):
    add_gradient_bar(slide, Emu(0), Emu(0), Inches(SLIDE_W_IN), Inches(0.09), GRADIENT_LEFT, GRADIENT_RIGHT)


def render_header(slide, insight, agent_label):
    x = Inches(MARGIN_X_IN)
    y = Inches(0.30)
    badge_w, badge_h = Inches(0.62), Inches(0.42)
    add_rect(slide, x, y, badge_w, badge_h, fill=C_BLUE_LIGHT_BG, radius=0.22)
    add_text(slide, x, y, badge_w, badge_h, f"{insight.number:02d}" if insight.number else "",
              16, C_BLUE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    div_x = x + badge_w + Inches(0.16)
    add_rect(slide, div_x, y + Inches(0.02), Inches(0.014), badge_h - Inches(0.04), fill=C_BORDER_LIGHT)

    tag_font = 11
    tag_w = Inches(pill_w(agent_label, tag_font, pad_in=0.36, bold=False))
    tag_h = Inches(0.4)
    tag_x = Inches(SLIDE_W_IN - MARGIN_X_IN) - tag_w
    add_pill(slide, tag_x, y + (badge_h - tag_h) / 2, tag_w, tag_h, agent_label,
              C_WHITE, C_BLUE, font_pt=tag_font, bold=False, border=C_BLUE)

    title_x = div_x + Inches(0.18)
    title_w = tag_x - Inches(0.2) - title_x
    title_font = 24
    title_h_in = text_height(insight.title, title_w / 914400, title_font, bold=True, line_mult=1.1)
    row_h = max(badge_h, Inches(title_h_in))
    add_text(slide, title_x, y, title_w, row_h, insight.title, title_font, C_NAVY,
              bold=True, anchor=MSO_ANCHOR.MIDDLE, line_mult=1.1)

    return y + row_h


def render_summary(slide, insight, y):
    x = Inches(MARGIN_X_IN)
    w = Inches(SLIDE_W_IN - 2 * MARGIN_X_IN)
    pad = Inches(0.16)
    font_pt = 12
    text_w_in = (w - 2 * pad - Inches(0.14)) / 914400
    h = Inches(text_height(insight.summary, text_w_in, font_pt, line_mult=1.3)) + 2 * pad
    add_rect(slide, x, y, w, h, fill=C_SUMMARY_BG, radius=0.06)
    add_rect(slide, x + Inches(0.02), y + Inches(0.1), Inches(0.045), h - Inches(0.2), fill=C_BLUE)
    add_text(slide, x + Inches(0.24), y + pad, w - Inches(0.24) - pad, h - 2 * pad,
              insight.summary, font_pt, C_MUTED_NAVY, line_mult=1.3, anchor=MSO_ANCHOR.MIDDLE)
    return y + h


def _bottom_row_height(insight, run_id, store, source_names, col_w_in):
    pad_in = 0.18
    label_h_in = 0.28
    wtm_h = 0.0
    if insight.interpretation:
        wtm_h = 2 * pad_in + label_h_in + text_height(
            insight.interpretation, col_w_in - 2 * pad_in, 11, line_mult=1.3)
    src_h = 0.0
    if insight.source_ids:
        names = [resolve_source_name(store, run_id, sid, source_names) for sid in insight.source_ids]
        row_x_in, rows, row_h_in = 0.0, 1, 0.32
        for name in names:
            pw_in = pill_w(name, 9.5, pad_in=0.28, bold=False)
            if row_x_in != 0.0 and (row_x_in + pw_in) > col_w_in:
                row_x_in = 0.0
                rows += 1
            row_x_in += pw_in + 0.12
        src_h = label_h_in + 0.12 + rows * row_h_in + (rows - 1) * 0.1
    return max(wtm_h, src_h, 1.1)


def render_bottom_row(slide, insight, run_id, store, source_names, y, row_h):
    x = Inches(MARGIN_X_IN)
    total_w = Inches(SLIDE_W_IN - 2 * MARGIN_X_IN)
    gap = Inches(0.3)
    wtm_w = (total_w - gap) * 0.46
    src_w = total_w - gap - wtm_w
    src_x = x + wtm_w + gap

    if insight.interpretation:
        add_rect(slide, x, y, wtm_w, row_h, fill=C_BLUE_LIGHT_BG, radius=0.04)
        add_rect(slide, x + Inches(0.02), y + Inches(0.12), Inches(0.045), row_h - Inches(0.24), fill=C_BLUE)
        pad = Inches(0.2)
        add_text(slide, x + pad, y + Inches(0.14), wtm_w - 2 * pad, Inches(0.26),
                  "WHAT THIS MEANS", 11, C_BLUE, bold=True)
        add_text(slide, x + pad, y + Inches(0.42), wtm_w - 2 * pad, row_h - Inches(0.56),
                  insight.interpretation, 11, C_MUTED_NAVY, line_mult=1.3)

    if insight.source_ids:
        add_text(slide, src_x, y + Inches(0.02), Inches(1.5), Inches(0.24), "SOURCES", 12, C_BLUE, bold=True)
        add_rect(slide, src_x, y + Inches(0.28), Inches(0.55), Inches(0.022), fill=C_BLUE)
        names = [resolve_source_name(store, run_id, sid, source_names) for sid in insight.source_ids]
        row_x, row_y, row_h_pill = src_x, y + Inches(0.42), Inches(0.32)
        for name in names:
            pw = Inches(pill_w(name, 9.5, pad_in=0.28, bold=False))
            if row_x != src_x and (row_x - src_x + pw) > src_w:
                row_x = src_x
                row_y += row_h_pill + Inches(0.1)
            add_pill(slide, row_x, row_y, pw, row_h_pill, name, C_BLUE_LIGHT_BG, C_BLUE,
                      font_pt=9.5, bold=False, border=RGBColor(0xC7, 0xDA, 0xFB))
            row_x += pw + Inches(0.12)
    elif not insight.interpretation:
        add_text(slide, src_x, y, src_w, Inches(0.3), "No source returned usable evidence for this card.",
                  10, C_TABLE_FAINT, italic=True)


def render_footer(slide, footer_text="\u00a9 2025 ProcDNA", page_num=None):
    y = Inches(SLIDE_H_IN - 0.32)
    add_text(slide, Inches(MARGIN_X_IN), y, Inches(3), Inches(0.22), footer_text, 9, C_FOOTER_GRAY)
    if page_num is not None:
        add_text(slide, Inches(SLIDE_W_IN - MARGIN_X_IN - 0.6), y, Inches(0.6), Inches(0.22),
                  str(page_num), 9, C_FOOTER_GRAY, align=PP_ALIGN.RIGHT)


# --------------------------------------------------------------------------- #
# Evidence renderers — each returns/measures a height that must fit inside a
# given `avail_h`, at a font/padding `scale` found by binary search.
# --------------------------------------------------------------------------- #
def _metrics_split(evidence):
    if len(evidence) > 4 and len(evidence) % 4 == 1:
        return evidence[:-1], evidence[-1]
    return evidence, None


def _metric_box_height(m, box_w_in, scale):
    val, lbl = clean_text(m.get("value", "")), clean_text(m.get("label", ""))
    pad = 0.14 * scale
    icon_d = 0.5 * scale
    inner_w = box_w_in - 2 * pad - icon_d - 0.1 * scale
    value_h = text_height(val, inner_w, 15 * scale, bold=True, line_mult=1.15)
    top_h = max(icon_d, value_h)
    label_h = text_height(lbl, box_w_in - 2 * pad, 9.5 * scale, line_mult=1.2)
    return pad + top_h + 0.08 * scale + 0.035 * scale + 0.06 * scale + label_h + pad


def _highlight_box_height(m, w_in, scale):
    val, lbl = clean_text(m.get("value", "")), clean_text(m.get("label", ""))
    pad = 0.16 * scale
    icon_d = 0.62 * scale
    inner_w = w_in - 2 * pad - icon_d - 0.16 * scale
    value_h = text_height(val, inner_w, 15 * scale, bold=True, line_mult=1.2)
    label_h = text_height(lbl, inner_w, 9.5 * scale, line_mult=1.2)
    return max(icon_d, value_h + 0.05 * scale + label_h) + 2 * pad


def metrics_height(evidence, w_in, scale):
    main, hi = _metrics_split(evidence)
    gap = 0.16 * scale
    total = 0.0
    if main:
        cols = 4
        box_w = (w_in - gap * (cols - 1)) / cols
        rows = [main[i:i + cols] for i in range(0, len(main), cols)]
        for row in rows:
            total += max(_metric_box_height(m, box_w, scale) for m in row) + gap
    if hi:
        total += _highlight_box_height(hi, w_in, scale) + gap
    return max(total - gap, 0.0)


def render_metrics(slide, evidence, x, y, w, scale):
    main, hi = _metrics_split(evidence)
    w_in = w / 914400
    gap = Inches(0.16 * scale)
    cursor_y = y
    if main:
        cols = 4
        box_w_in = (w_in - (0.16 * scale) * (cols - 1)) / cols
        box_w = Inches(box_w_in)
        rows = [main[i:i + cols] for i in range(0, len(main), cols)]
        for row in rows:
            row_h = Inches(max(_metric_box_height(m, box_w_in, scale) for m in row))
            cursor_x = x
            for i, m in enumerate(row):
                add_rect(slide, cursor_x, cursor_y, box_w, row_h, fill=C_WHITE, line=C_BORDER_LIGHT, radius=0.06)
                pad = Inches(0.14 * scale)
                icon_d = Inches(0.5 * scale)
                bg, fg = C_ICON_BG[i % len(C_ICON_BG)]
                add_shape(slide, MSO_SHAPE.OVAL, cursor_x + pad, cursor_y + pad, icon_d, icon_d, fill=bg)
                icon_kind = pick_icon(m.get("label", ""), i)
                draw_icon(slide, icon_kind, cursor_x + pad + icon_d * 0.2, cursor_y + pad + icon_d * 0.2,
                          icon_d * 0.6, fg)
                val, lbl = clean_text(m.get("value", "")), clean_text(m.get("label", ""))
                text_x = cursor_x + pad + icon_d + Inches(0.1 * scale)
                text_w = box_w - pad - (text_x - cursor_x)
                value_font = 15 * scale
                value_h = Inches(text_height(val, text_w / 914400, value_font, bold=True, line_mult=1.15))
                add_text(slide, text_x, cursor_y + pad, text_w, value_h, val, value_font, C_NAVY,
                          bold=True, line_mult=1.15)
                top_h = max(icon_d, value_h)
                underline_y = cursor_y + pad + top_h + Inches(0.08 * scale)
                add_rect(slide, cursor_x + pad, underline_y, Inches(0.35 * scale), Inches(0.035 * scale), fill=fg)
                label_y = underline_y + Inches(0.035 * scale) + Inches(0.06 * scale)
                add_text(slide, cursor_x + pad, label_y, box_w - 2 * pad, row_h - (label_y - cursor_y) - pad,
                          lbl, 9.5 * scale, C_TABLE_FAINT, line_mult=1.2)
                cursor_x += box_w + gap
            cursor_y += row_h + gap
    if hi:
        h = Inches(_highlight_box_height(hi, w_in, scale))
        add_rect(slide, x, cursor_y, w, h, fill=C_WHITE, line=C_BORDER_LIGHT, radius=0.06)
        pad = Inches(0.16 * scale)
        icon_d = Inches(0.62 * scale)
        bg, fg = C_ICON_BG[0]
        add_shape(slide, MSO_SHAPE.OVAL, x + pad, cursor_y + (h - icon_d) / 2, icon_d, icon_d, fill=bg)
        draw_icon(slide, pick_icon(hi.get("label", ""), 0), x + pad + icon_d * 0.2,
                  cursor_y + (h - icon_d) / 2 + icon_d * 0.2, icon_d * 0.6, fg)
        val, lbl = clean_text(hi.get("value", "")), clean_text(hi.get("label", ""))
        text_x = x + pad + icon_d + Inches(0.16 * scale)
        text_w = w - pad - (text_x - x)
        value_font = 15 * scale
        value_h = Inches(text_height(val, text_w / 914400, value_font, bold=True, line_mult=1.2))
        block_h = value_h + Inches(0.05 * scale) + Inches(text_height(lbl, text_w / 914400, 9.5 * scale, line_mult=1.2))
        block_y = cursor_y + (h - block_h) / 2
        add_text(slide, text_x, block_y, text_w, value_h, val, value_font, C_NAVY, bold=True, line_mult=1.2)
        add_text(slide, text_x, block_y + value_h + Inches(0.05 * scale), text_w, h,
                  lbl, 9.5 * scale, C_TABLE_FAINT, line_mult=1.2)
        cursor_y += h + gap
    return cursor_y - gap


def _col_weights(columns, rows):
    weights = []
    for i, col in enumerate(columns):
        longest = len(col)
        for r in rows:
            if i < len(r):
                longest = max(longest, len(clean_text(r[i])))
        weights.append(max(longest, 8))
    total = sum(weights)
    return [w_ / total for w_ in weights]


def table_height(evidence, w_in, scale):
    cols, rows = evidence.get("columns", []), evidence.get("rows", [])
    if not cols or not rows:
        return 0.0
    weights = _col_weights(cols, rows)
    col_w_in = [w_in * wt for wt in weights]
    pad_in = 0.12 * scale
    cell_font = 10.5 * scale
    total = 0.34 * scale
    for row in rows:
        cell_texts = [clean_text(v) for v in row]
        total += max(
            text_height(t, col_w_in[i] - 2 * pad_in, cell_font, line_mult=1.2)
            for i, t in enumerate(cell_texts)
        ) + 2 * pad_in
    return total


def render_table(slide, evidence, x, y, w, scale):
    cols, rows = evidence.get("columns", []), evidence.get("rows", [])
    if not cols or not rows:
        return y
    w_in = w / 914400
    weights = _col_weights(cols, rows)
    col_w_in = [w_in * wt for wt in weights]
    pad_in = 0.12 * scale
    header_font, cell_font = 9.5 * scale, 10.5 * scale

    header_h = Inches(0.34 * scale)
    add_rect(slide, x, y, w, header_h, fill=C_TABLE_HEADER_BG)
    cx = x
    for c, cw in zip(cols, col_w_in):
        add_text(slide, cx + Inches(pad_in), y, Inches(cw - 2 * pad_in), header_h, c.upper(),
                  header_font, C_MUTED_NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        cx += Inches(cw)
    cursor_y = y + header_h

    for row in rows:
        cell_texts = [clean_text(v) for v in row]
        row_h = Inches(max(
            text_height(t, col_w_in[i] - 2 * pad_in, cell_font, line_mult=1.2)
            for i, t in enumerate(cell_texts)
        ) + 2 * pad_in)
        cx = x
        for i, t in enumerate(cell_texts):
            faint = t.strip().lower() in ("not quantified", "n/a", "unknown", "not reported")
            add_text(slide, cx + Inches(pad_in), cursor_y, Inches(col_w_in[i] - 2 * pad_in), row_h, t,
                      cell_font, C_TABLE_FAINT if faint else C_TABLE_TEXT, italic=faint,
                      anchor=MSO_ANCHOR.MIDDLE, line_mult=1.2)
            cx += Inches(col_w_in[i])
        cursor_y += row_h
        add_rect(slide, x, cursor_y - Inches(0.006), w, Inches(0.006), fill=C_BORDER_LIGHT)
    add_rect(slide, x, y, w, cursor_y - y, fill=None, line=C_BORDER_LIGHT, radius=0.015)
    return cursor_y


def _steps_grid(n, w_in, min_chevron_in=1.35):
    max_cols = max(1, min(n, int(w_in / min_chevron_in)))
    import math
    rows = math.ceil(n / max_cols)
    cols = math.ceil(n / rows)
    return cols, rows


def render_process_summary(slide, insight, y):
    """
    Larger summary treatment used ONLY on process / steps slides.
    """
    x = Inches(MARGIN_X_IN)
    w = Inches(SLIDE_W_IN - 2 * MARGIN_X_IN)

    pad_x = Inches(0.22)
    pad_y = Inches(0.20)

    font_pt = 13.0

    text_w_in = (
        w
        - Inches(0.28)
        - pad_x
        - Inches(0.12)
    ) / 914400

    text_h_in = text_height(
        insight.summary,
        text_w_in,
        font_pt,
        line_mult=1.30,
    )

    # Keep the summary visually substantial like the reference slide.
    h = max(
        Inches(0.78),
        Inches(text_h_in) + 2 * pad_y,
    )

    add_rect(
        slide,
        x,
        y,
        w,
        h,
        fill=C_SUMMARY_BG,
        radius=0.06,
    )

    # Blue vertical accent.
    add_rect(
        slide,
        x + Inches(0.02),
        y + Inches(0.10),
        Inches(0.045),
        h - Inches(0.20),
        fill=C_BLUE,
    )

    add_text(
        slide,
        x + Inches(0.24),
        y + pad_y,
        w - Inches(0.24) - pad_x,
        h - 2 * pad_y,
        insight.summary,
        font_pt,
        C_MUTED_NAVY,
        line_mult=1.30,
        anchor=MSO_ANCHOR.MIDDLE,
    )

    return y + h


# --------------------------------------------------------------------------- #
# PROCESS-ONLY BOTTOM ROW
# --------------------------------------------------------------------------- #
def render_process_bottom_row(
    slide,
    insight,
    run_id,
    store,
    source_names,
    y,
    row_h,
):
    """
    Bottom section matching the reference process slide.

    Used ONLY for evidence_type == "steps".
    """

    x = Inches(MARGIN_X_IN)
    total_w = Inches(SLIDE_W_IN - 2 * MARGIN_X_IN)

    # Reference has a wider WHAT THIS MEANS card.
    gap = Inches(0.14)
    wtm_w = Inches(6.42)
    src_w = total_w - wtm_w - gap
    src_x = x + wtm_w + gap

    # ------------------------------------------------------------------ #
    # WHAT THIS MEANS
    # ------------------------------------------------------------------ #
    if insight.interpretation:
        add_rect(
            slide,
            x,
            y,
            wtm_w,
            row_h,
            fill=C_BLUE_LIGHT_BG,
            radius=0.045,
        )

        add_rect(
            slide,
            x + Inches(0.02),
            y + Inches(0.15),
            Inches(0.045),
            row_h - Inches(0.30),
            fill=C_BLUE,
        )

        pad_x = Inches(0.22)

        add_text(
            slide,
            x + pad_x,
            y + Inches(0.17),
            wtm_w - 2 * pad_x,
            Inches(0.25),
            "WHAT THIS MEANS",
            11.5,
            C_BLUE,
            bold=True,
        )

        add_text(
            slide,
            x + pad_x,
            y + Inches(0.51),
            wtm_w - 2 * pad_x,
            row_h - Inches(0.66),
            insight.interpretation,
            11.0,
            C_MUTED_NAVY,
            line_mult=1.32,
        )

    # ------------------------------------------------------------------ #
    # SOURCES CARD
    # ------------------------------------------------------------------ #
    if insight.source_ids:
        add_rect(
            slide,
            src_x,
            y,
            src_w,
            row_h,
            fill=C_WHITE,
            line=C_BORDER_LIGHT,
            radius=0.045,
        )

        add_text(
            slide,
            src_x + Inches(0.22),
            y + Inches(0.19),
            Inches(1.6),
            Inches(0.25),
            "SOURCES",
            12,
            C_BLUE,
            bold=True,
        )

        add_rect(
            slide,
            src_x + Inches(0.22),
            y + Inches(0.48),
            Inches(0.42),
            Inches(0.024),
            fill=C_BLUE,
        )

        names = [
            resolve_source_name(
                store,
                run_id,
                sid,
                source_names,
            )
            for sid in insight.source_ids
        ]

        row_x = src_x + Inches(0.22)
        row_y = y + Inches(0.68)
        pill_h = Inches(0.34)
        right_limit = src_x + src_w - Inches(0.22)

        for name in names:
            pw = Inches(
                pill_w(
                    name,
                    9.5,
                    pad_in=0.30,
                    bold=False,
                )
            )

            if (
                row_x > src_x + Inches(0.22)
                and row_x + pw > right_limit
            ):
                row_x = src_x + Inches(0.22)
                row_y += pill_h + Inches(0.11)

            add_pill(
                slide,
                row_x,
                row_y,
                pw,
                pill_h,
                name,
                C_BLUE_LIGHT_BG,
                C_BLUE,
                font_pt=9.5,
                bold=False,
                border=RGBColor(0xC7, 0xDA, 0xFB),
            )

            row_x += pw + Inches(0.12)


# --------------------------------------------------------------------------- #
# PROCESS STAGE HEIGHT
# --------------------------------------------------------------------------- #
def _chevron_content_height(text, chevron_w_in, scale):
    """
    Height calculation ONLY for process / steps cards.
    """

    # Do not allow the generic scale-up to make process elements huge.
    s = min(scale, 1.0)

    # Much wider text block than the previous version.
    # This prevents:
    #   Symptomatic / clinical / presentation
    # and instead gives:
    #   Symptomatic clinical
    #   presentation
    label_w_in = chevron_w_in * 0.84

    label_font = 10.5 * s

    lines = wrap_text(
        clean_text(text),
        label_w_in,
        label_font,
        bold=True,
    )

    label_h_in = (
        len(lines)
        * label_font
        * 1.15
        / 72.0
    )

    icon_d = 0.42 * s
    icon_to_text_gap = 0.10 * s

    # Space inside the arrow above the icon.
    top_inside = 0.43 * s
    bottom_pad = 0.17 * s

    calculated = (
        top_inside
        + icon_d
        + icon_to_text_gap
        + label_h_in
        + bottom_pad
    )

    # Consistent arrow height.
    return max(
        1.62 * s,
        calculated,
    )


# --------------------------------------------------------------------------- #
# PROCESS PANEL HEIGHT
# --------------------------------------------------------------------------- #
def steps_height(evidence, w_in, scale):
    """
    Total height of process panel ONLY.
    """

    n = len(evidence)

    if n == 0:
        return 0.0

    s = min(scale, 1.0)

    cols, rows_n = _steps_grid(
        n,
        w_in,
    )

    panel_pad_x = 0.12 * s

    # Narrow connectors allow wider stage cards.
    sep_w_in = 0.12 * s

    inner_w_in = (
        w_in
        - 2 * panel_pad_x
    )

    chevron_w_in = (
        inner_w_in
        - sep_w_in * (cols - 1)
    ) / cols

    # Important:
    # Larger top padding moves the number badges down within the
    # outer card just like the reference slide.
    panel_top_pad = 0.50 * s
    panel_bottom_pad = 0.20 * s
    row_gap = 0.22 * s

    total = panel_top_pad

    for r in range(rows_n):
        row_items = evidence[
            r * cols:(r + 1) * cols
        ]

        row_h = max(
            _chevron_content_height(
                step,
                chevron_w_in,
                s,
            )
            for step in row_items
        )

        total += row_h

        if r < rows_n - 1:
            total += row_gap

    total += panel_bottom_pad

    return total


# --------------------------------------------------------------------------- #
# PROCESS RENDERER
# --------------------------------------------------------------------------- #
def render_steps(slide, evidence, x, y, w, scale):
    """
    Renderer ONLY for evidence_type == "steps".

    Fixes:
      - journey sits lower on slide
      - larger surrounding process panel
      - badges no longer sit too close to panel top
      - wider text inside each stage
      - smaller gaps between stages
      - chevrons fill more horizontal space
      - better icon / label positioning
    """

    n = len(evidence)

    if n == 0:
        return y

    s = min(scale, 1.0)

    w_in = w / 914400

    cols, rows_n = _steps_grid(
        n,
        w_in,
    )

    # ------------------------------------------------------------------ #
    # GEOMETRY
    # ------------------------------------------------------------------ #

    panel_pad_x = Inches(0.12 * s)

    # Extra top breathing room is intentional.
    panel_top_pad = Inches(0.50 * s)
    panel_bottom_pad = Inches(0.20 * s)

    sep_w = Inches(0.12 * s)
    row_gap = Inches(0.22 * s)

    inner_x = x + panel_pad_x
    inner_w = w - 2 * panel_pad_x

    chevron_w = (
        inner_w
        - sep_w * (cols - 1)
    ) / cols

    chevron_w_in = (
        chevron_w / 914400
    )

    row_heights = []

    for r in range(rows_n):

        row_items = evidence[
            r * cols:(r + 1) * cols
        ]

        row_h_in = max(
            _chevron_content_height(
                step,
                chevron_w_in,
                s,
            )
            for step in row_items
        )

        row_heights.append(
            Inches(row_h_in)
        )

    panel_h = (
        panel_top_pad
        + sum(row_heights)
        + row_gap * max(0, rows_n - 1)
        + panel_bottom_pad
    )

    # ------------------------------------------------------------------ #
    # OUTER JOURNEY CONTAINER
    # ------------------------------------------------------------------ #

    add_rect(
        slide,
        x,
        y,
        w,
        panel_h,
        fill=C_WHITE,
        line=C_BORDER_LIGHT,
        radius=0.04,
    )

    cursor_y = (
        y
        + panel_top_pad
    )

    # ------------------------------------------------------------------ #
    # PROCESS STEPS
    # ------------------------------------------------------------------ #

    for r in range(rows_n):

        row_items = evidence[
            r * cols:(r + 1) * cols
        ]

        row_h = row_heights[r]
        cursor_x = inner_x

        for i, step in enumerate(row_items):

            idx = (
                r * cols
                + i
            )

            text = clean_text(step)

            fill = C_PASTEL[
                idx % len(C_PASTEL)
            ]

            # ---------------------------------------------------------- #
            # CHEVRON
            # ---------------------------------------------------------- #

            add_shape(
                slide,
                MSO_SHAPE.CHEVRON,
                cursor_x,
                cursor_y,
                chevron_w,
                row_h,
                fill=fill,

                # Flatter arrow than the previous 0.32.
                adjustments=[0.20],
            )

            stage_center_x = (
                cursor_x
                + chevron_w / 2
            )

            # ---------------------------------------------------------- #
            # NUMBER BADGE
            # ---------------------------------------------------------- #

            badge_d = Inches(
                0.38 * s
            )

            badge_x = (
                stage_center_x
                - badge_d / 2
            )

            badge_y = (
                cursor_y
                - badge_d / 2
            )

            badge_color = (
                C_TEAL
                if (idx + 1) % 3 == 2
                else C_BLUE
            )

            add_shape(
                slide,
                MSO_SHAPE.OVAL,
                badge_x,
                badge_y,
                badge_d,
                badge_d,
                fill=badge_color,
            )

            add_text(
                slide,
                badge_x,
                badge_y,
                badge_d,
                badge_d,
                str(idx + 1),
                10,
                C_WHITE,
                bold=True,
                align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE,
            )

            # ---------------------------------------------------------- #
            # ICON
            # ---------------------------------------------------------- #

            icon_d = Inches(
                0.42 * s
            )

            icon_x = (
                stage_center_x
                - icon_d / 2
            )

            icon_y = (
                cursor_y
                + Inches(0.43 * s)
            )

            icon_color = (
                C_TEAL
                if (idx + 1) % 3 == 2
                else C_BLUE
            )

            draw_icon(
                slide,
                pick_icon(
                    text,
                    idx,
                ),
                icon_x,
                icon_y,
                icon_d,
                icon_color,
            )

            # ---------------------------------------------------------- #
            # TEXT
            # ---------------------------------------------------------- #

            # Wider label is the important change here.
            label_w = (
                chevron_w * 0.84
            )

            label_x = (
                stage_center_x
                - label_w / 2
            )

            label_y = (
                icon_y
                + icon_d
                + Inches(0.10 * s)
            )

            label_h = (
                cursor_y
                + row_h
                - Inches(0.11)
                - label_y
            )

            add_text(
                slide,
                label_x,
                label_y,
                label_w,
                label_h,
                text,
                10.5 * s,
                C_NAVY,
                bold=True,
                align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.TOP,
                line_mult=1.15,
            )

            # ---------------------------------------------------------- #
            # CONNECTOR
            # ---------------------------------------------------------- #

            if i < len(row_items) - 1:

                sep_x = (
                    cursor_x
                    + chevron_w
                )

                add_text(
                    slide,
                    sep_x,
                    cursor_y
                    + row_h / 2
                    - Inches(0.13),
                    sep_w,
                    Inches(0.26),
                    "\u203a",
                    17,
                    C_TABLE_FAINT,
                    bold=True,
                    align=PP_ALIGN.CENTER,
                    anchor=MSO_ANCHOR.MIDDLE,
                )

            cursor_x += (
                chevron_w
                + sep_w
            )

        cursor_y += row_h

        if r < rows_n - 1:
            cursor_y += row_gap

    return y + panel_h





def _synth_cell_height(item, col_w_in, scale):
    badge_d = 0.34 * scale
    pad = 0.14 * scale
    text_w = col_w_in - 2 * pad - badge_d - 0.12 * scale
    text_h = text_height(clean_text(item), text_w, 10.5 * scale, line_mult=1.28)
    return max(badge_d, text_h) + 2 * pad


def synth_height(evidence, w_in, scale):
    n = len(evidence)
    if n == 0:
        return 0.0
    divider_gap = 0.3 * scale
    col_w_in = (w_in - divider_gap) / 2
    row_gap = 0.14 * scale
    rows = (n + 1) // 2
    total = 0.0
    for r in range(rows):
        left = evidence[2 * r]
        right = evidence[2 * r + 1] if 2 * r + 1 < n else None
        left_h = _synth_cell_height(left, col_w_in, scale)
        right_h = _synth_cell_height(right, col_w_in, scale) if right else 0.0
        total += max(left_h, right_h) + row_gap
    return total - row_gap


def render_synth(slide, evidence, x, y, w, scale):
    n = len(evidence)
    if n == 0:
        return y
    w_in = w / 914400
    divider_gap = Inches(0.3 * scale)
    col_w = (w - divider_gap) / 2
    col_w_in = col_w / 914400
    row_gap = Inches(0.14 * scale)
    rows = (n + 1) // 2

    def draw_cell(item, cell_x, cell_y, cell_h):
        add_rect(slide, cell_x, cell_y, col_w, cell_h, fill=C_BLUE_LIGHT_BG, radius=0.08)
        badge_d = Inches(0.34 * scale)
        pad = Inches(0.14 * scale)
        add_shape(slide, MSO_SHAPE.OVAL, cell_x + pad, cell_y + pad, badge_d, badge_d, fill=C_BLUE)
        num = evidence.index(item) + 1
        add_text(slide, cell_x + pad, cell_y + pad, badge_d, badge_d, str(num), 10 * scale, C_WHITE,
                  bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text_x = cell_x + pad + badge_d + Inches(0.12 * scale)
        text_w = cell_x + col_w - pad - text_x
        add_text(slide, text_x, cell_y + pad, text_w, cell_h - 2 * pad, clean_text(item),
                  10.5 * scale, C_NAVY, line_mult=1.28, anchor=MSO_ANCHOR.MIDDLE)

    cursor_y = y
    for r in range(rows):
        left = evidence[2 * r]
        right = evidence[2 * r + 1] if 2 * r + 1 < n else None
        left_h = _synth_cell_height(left, col_w_in, scale)
        right_h = _synth_cell_height(right, col_w_in, scale) if right else 0.0
        row_h = Inches(max(left_h, right_h))
        draw_cell(left, x, cursor_y, row_h)
        if right:
            draw_cell(right, x + col_w + divider_gap, cursor_y, row_h)
        cursor_y += row_h + row_gap
    bottom = cursor_y - row_gap
    if n > 1:
        add_rect(slide, x + col_w + divider_gap / 2, y, Inches(0.012), bottom - y, fill=C_BORDER_LIGHT)
    return bottom


# --------------------------------------------------------------------------- #
# Dispatch + shrink-to-fit
# --------------------------------------------------------------------------- #
def _evidence_height(insight, w_in, scale):
    if not insight.evidence:
        return 0.0
    et = insight.evidence_type
    if et == "metrics":
        return metrics_height(insight.evidence, w_in, scale)
    if et == "table" and "columns" in insight.evidence:
        return table_height(insight.evidence, w_in, scale)
    if et == "steps":
        return steps_height(insight.evidence, w_in, scale)
    return synth_height(insight.evidence, w_in, scale)


def _render_evidence(slide, insight, x, y, w, scale):
    et = insight.evidence_type
    if et == "metrics":
        return render_metrics(slide, insight.evidence, x, y, w, scale)
    if et == "table" and "columns" in insight.evidence:
        return render_table(slide, insight.evidence, x, y, w, scale)
    if et == "steps":
        return render_steps(slide, insight.evidence, x, y, w, scale)
    return render_synth(slide, insight.evidence, x, y, w, scale)


def _fit_scale(insight, content_w_in, available_h_in, iters=16):
    """Binary search for the LARGEST scale in [MIN_SCALE, MAX_SCALE] whose
    measured evidence height still fits inside `available_h_in`. Unlike a
    simple "shrink if too big" check, this also scales UP short content
    (e.g. a 4-step journey, a 5-metric card) so it fills the same reserved
    band the template always allocates, instead of leaving a dead gap above
    the WHAT THIS MEANS / SOURCES row."""
    if not insight.evidence:
        return 1.0
    lo, hi = MIN_SCALE, MAX_SCALE
    lo_h = _evidence_height(insight, content_w_in, lo)
    if lo_h > available_h_in:
        return lo
    hi_h = _evidence_height(insight, content_w_in, hi)
    if hi_h <= available_h_in:
        return hi
    for _ in range(iters):
        mid = (lo + hi) / 2
        mid_h = _evidence_height(insight, content_w_in, mid)
        if mid_h <= available_h_in:
            lo = mid
        else:
            hi = mid
    return lo


# --------------------------------------------------------------------------- #
# Full slide assembly
# --------------------------------------------------------------------------- #
def build_insight_pptx(
    insight,
    run_id,
    store,
    source_names=None,
    footer_text="\u00a9 2025 ProcDNA",
):
    prs = Presentation()

    prs.slide_width = Inches(
        SLIDE_W_IN
    )

    prs.slide_height = Inches(
        SLIDE_H_IN
    )

    slide = prs.slides.add_slide(
        prs.slide_layouts[6]
    )

    slide.background.fill.solid()

    slide.background.fill.fore_color.rgb = (
        C_WHITE
    )

    # ================================================================== #
    # PROCESS / STEPS SLIDE
    # ================================================================== #

    if insight.evidence_type == "steps":

        # Process slide-specific label.
        if insight.category == "Clinical":
            agent_label = "Clinical Landscape"
        else:
            agent_label = insight.category

        render_top_bar(
            slide
        )

        header_bottom = render_header(
            slide,
            insight,
            agent_label,
        )

        # Slightly more breathing room below title.
        summary_bottom = render_process_summary(
            slide,
            insight,
            header_bottom + Inches(0.24),
        )

        content_x = Inches(
            MARGIN_X_IN
        )

        content_w = Inches(
            SLIDE_W_IN
            - 2 * MARGIN_X_IN
        )

        content_w_in = (
            content_w / 914400
        )

        # -------------------------------------------------------------- #
        # IMPORTANT LAYOUT FIX
        #
        # Previously:
        # bottom cards were anchored near the footer (~5.6"),
        # creating the huge blank gap visible in your screenshot.
        #
        # Process slide now uses the reference geometry.
        # -------------------------------------------------------------- #

        bottom_row_y_in = 5.00
        bottom_row_h_in = 1.52

        # Do not allow process panel to start too high.
        content_y_in = max(
            summary_bottom / 914400 + 0.18,
            2.28,
        )

        available_h_in = (
            bottom_row_y_in
            - content_y_in
            - 0.20
        )

        # Keep process at normal visual size.
        # It can shrink only if absolutely necessary.
        scale = min(
            1.0,
            _fit_scale(
                insight,
                content_w_in,
                available_h_in,
            ),
        )

        if insight.evidence:

            _render_evidence(
                slide,
                insight,
                content_x,
                Inches(content_y_in),
                content_w,
                scale,
            )

        render_process_bottom_row(
            slide,
            insight,
            run_id,
            store,
            source_names,
            Inches(bottom_row_y_in),
            Inches(bottom_row_h_in),
        )

        render_footer(
            slide,
            footer_text=footer_text,
            page_num=1,
        )

    # ================================================================== #
    # ALL OTHER CARD TYPES
    #
    # THIS IS YOUR ORIGINAL LOGIC.
    # Metrics / table / synthesis are NOT changed.
    # ================================================================== #

    else:

        agent_label = (
            "Celestra Synthesis"
            if insight.category == "Synthesis"
            else insight.category
        )

        render_top_bar(
            slide
        )

        header_bottom = render_header(
            slide,
            insight,
            agent_label,
        )

        summary_bottom = render_summary(
            slide,
            insight,
            header_bottom + Inches(0.16),
        )

        content_x = Inches(
            MARGIN_X_IN
        )

        content_w = Inches(
            SLIDE_W_IN
            - 2 * MARGIN_X_IN
        )

        content_w_in = (
            content_w / 914400
        )

        col_w_in = (
            SLIDE_W_IN
            - 2 * MARGIN_X_IN
            - 0.3
        ) * 0.46

        bottom_row_h_in = _bottom_row_height(
            insight,
            run_id,
            store,
            source_names,
            col_w_in,
        )

        footer_top_in = (
            SLIDE_H_IN
            - 0.5
        )

        bottom_row_y_in = (
            footer_top_in
            - bottom_row_h_in
        )

        content_y_in = (
            summary_bottom / 914400
            + 0.18
        )

        available_h_in = max(
            0.6,
            bottom_row_y_in
            - 0.18
            - content_y_in,
        )

        scale = _fit_scale(
            insight,
            content_w_in,
            available_h_in,
        )

        if insight.evidence:

            _render_evidence(
                slide,
                insight,
                content_x,
                Inches(content_y_in),
                content_w,
                scale,
            )

        render_bottom_row(
            slide,
            insight,
            run_id,
            store,
            source_names,
            Inches(bottom_row_y_in),
            Inches(bottom_row_h_in),
        )

        render_footer(
            slide,
            footer_text=footer_text,
        )

    stream = io.BytesIO()

    prs.save(
        stream
    )

    stream.seek(0)

    return stream

def register_route(app, store):
    @app.get("/runs/{run_id}/insights/{insight_id}/ppt")
    async def export_insight_ppt(run_id: str, insight_id: str):
        insight = store.get_insight(run_id, insight_id)
        if not insight:
            raise HTTPException(404, "Insight not found")

        stream = build_insight_pptx(insight, run_id, store)

        safe_title = "".join(c if c.isalnum() else "_" for c in insight.title)
        filename = f"Insight_{safe_title}.pptx"

        return StreamingResponse(
            stream,
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )

# //claude working