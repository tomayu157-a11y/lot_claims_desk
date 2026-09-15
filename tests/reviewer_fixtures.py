"""Small files built in code for the reviewer-file tests, so no binary
fixtures live in the repository."""
from __future__ import annotations

import base64
import io

_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
)


def pdf_bytes(pages: int = 3, text: bool = True, password: str = "") -> bytes:
    """A PDF with a large-font heading and a body line per page."""
    import pymupdf

    doc = pymupdf.open()
    for n in range(pages):
        page = doc.new_page()
        if text:
            page.insert_text((72, 80), f"Section {n + 1} Heading", fontsize=20)
            page.insert_text((72, 120), f"Body paragraph on page {n + 1} about frontline regimens.",
                             fontsize=11)
    if password:
        return doc.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256,
                           user_pw=password, owner_pw=password + "-owner")
    return doc.tobytes()


def docx_bytes() -> bytes:
    """A Word document with two heading levels, a bullet, a table and an image."""
    from docx import Document
    from docx.shared import Inches

    d = Document()
    d.add_heading("Payer policy", 1)
    d.add_paragraph("Frontline venetoclax is covered for del(17p).")
    d.add_heading("Step therapy", 2)
    d.add_paragraph("Try BTK inhibitors first.", style="List Bullet")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Regimen", "Line"
    table.cell(1, 0).text, table.cell(1, 1).text = "VenO", "1L"
    d.add_picture(io.BytesIO(_PNG), width=Inches(1))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
